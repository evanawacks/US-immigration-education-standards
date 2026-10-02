#!/usr/bin/env python3
"""Extract plain text from every file in data/manifest.csv into data/text/<ABBR>/<name>.txt
and write data/text_stats.csv (words, pages, sha256 of text, extraction status)."""
import csv, hashlib, subprocess, re
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "text"

def pdf(p):
    r = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True)
    pages = r.stdout.count("\f")
    return r.stdout, pages

def docx_(p):
    import docx
    d = docx.Document(str(p)); parts = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            seen = []
            for c in row.cells:
                if c.text not in seen: seen.append(c.text)
            parts.append(" | ".join(seen))
    return "\n".join(parts), None

def xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True); parts = []
    for ws in wb.worksheets:
        parts.append(f"## sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            v = [str(c) for c in row if c is not None]
            if v: parts.append(" | ".join(v))
    return "\n".join(parts), len(wb.worksheets)

rows = list(csv.DictReader(open(ROOT / "data" / "manifest.csv")))
stats = []
for r in rows:
    p = ROOT / r["local_path"]; ext = r["file_type"]
    try:
        text, pages = {"pdf": pdf, "docx": docx_, "xlsx": xlsx}[ext](p)
        status = "ok"
    except Exception as e:
        text, pages, status = "", None, f"error {type(e).__name__}: {e}"
    words = len(text.split())
    if status == "ok" and words < 50: status = "no-text (scanned? needs OCR)"
    d = OUT / r["abbr"]; d.mkdir(parents=True, exist_ok=True)
    (d / (p.name + ".txt")).write_text(text)
    stats.append({"abbr": r["abbr"], "file": p.name, "local_path": r["local_path"], "file_type": ext, "pages_or_sheets": pages,
                  "words": words, "file_sha256": r["sha256"], "text_sha256": hashlib.sha256(" ".join(text.split()).encode()).hexdigest(), "status": status})
with open(ROOT / "data" / "text_stats.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=stats[0].keys()); w.writeheader(); w.writerows(stats)
print(len(stats), "files;", sum(s["words"] for s in stats), "words;", sum(s["status"] != "ok" for s in stats), "problems")
