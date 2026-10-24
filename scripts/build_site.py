#!/usr/bin/env python3
"""Build the data files for the review website (index.html + site/).

For every state it writes, under site/data/:
  index.json                 states, their files (used / not used and why), and the database rows (no text)
  rows/<ST>.json             the full database rows, each column split into the units it was built from
  docs/<ST>/<n>.json         one per used document: its units (clean text) with label, page position, and
                             the database rows each unit was copied into
  pdf/<ST>/<n>.pdf           .docx sources rendered to PDF with LibreOffice (so they can be shown as pages)
  cache/<ST>/<n>.bbox.json   cached unit positions (rebuilt when the source file changes)

Usage: python scripts/build_site.py [ST ...]    (needs pymupdf; LibreOffice `soffice` for .docx)
       python scripts/build_site.py --topics     only refresh site/data/topics/, the tags in each document and the main
                                                 documents / subfiles (after topic_tags.py or keyword_tags.py; no pymupdf)
"""
import csv, hashlib, json, re, subprocess, sys, shutil, tempfile
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from assemble_clean import assemble_state, load_cfg, CLEAN
from check_labels import parse_units
import doc_parts
OUT = ROOT / "site" / "data"

STATE_NAMES = dict(AL="Alabama", AK="Alaska", AZ="Arizona", AR="Arkansas", CA="California", CO="Colorado", CT="Connecticut", DE="Delaware",
    DC="District of Columbia", FL="Florida", GA="Georgia", HI="Hawaii", ID="Idaho", IL="Illinois", IN="Indiana", IA="Iowa", KS="Kansas",
    KY="Kentucky", LA="Louisiana", ME="Maine", MD="Maryland", MA="Massachusetts", MI="Michigan", MN="Minnesota", MS="Mississippi",
    MO="Missouri", MT="Montana", NE="Nebraska", NV="Nevada", NH="New Hampshire", NJ="New Jersey", NM="New Mexico", NY="New York",
    NC="North Carolina", ND="North Dakota", OH="Ohio", OK="Oklahoma", OR="Oregon", PA="Pennsylvania", RI="Rhode Island",
    SC="South Carolina", SD="South Dakota", TN="Tennessee", TX="Texas", UT="Utah", VT="Vermont", VA="Virginia", WA="Washington",
    WV="West Virginia", WI="Wisconsin", WY="Wyoming")
COLS = ["body_standards", "body_examples", "backup"]
NO_RENDER = {"MO curr-mls-standards-ss-6-12-sboe-2016.docx"}   # LibreOffice hangs on these (>25 min)

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def dump(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")))

def unit_tag(u):
    if u["kind"] == "cell": return f"T{u['table']} R{u['row']} C{u['col']}"
    if u["kind"] == "row": return f"R{u['row']}"
    if u["kind"] == "sheet": return "sheet"
    return u.get("style") or ""

# ---------- unit positions on the page ----------
def pdf_positions(src, units):
    """Re-run the PDF extraction with positions; use them only if it reproduces the stored units exactly."""
    from extract_units import pdf_units, ocr_files
    again, _ = pdf_units(src, with_bbox=True, ocr=src.name in ocr_files(src.parent.name))
    if [(u["page"], u["text"]) for u in again] != [(u["page"], u["text"]) for u in units]:
        return None, "re-extraction did not reproduce the stored units; no page highlights"
    return {i + 1: [u["page"], u["bbox"]] for i, u in enumerate(again)}, None

def docx_positions(pdf, units):
    """Find each unit of a .docx in its LibreOffice rendering by searching for the start and end of its text."""
    import pymupdf
    doc = pymupdf.open(str(pdf))
    words = lambda t, n, end=False: " ".join((t.split()[-n:] if end else t.split()[:n]))
    pos, cur, miss = {}, 0, 0
    for u in units:
        t = u["text"]
        if u.get("furniture") or not t.strip(): continue
        n = int(u["id"][1:]); found = None
        order = list(range(cur, min(cur + 4, len(doc)))) + list(range(0, len(doc)))
        for nw in (8, 4):
            for pi in order:
                hits = doc[pi].search_for(words(t, nw))
                if hits:
                    found = (pi, hits); break
            if found: break
        if not found:
            miss += 1; continue
        pi, hits = found
        page = doc[pi]; W, H = page.rect.width, page.rect.height
        r = pymupdf.Rect(hits[0])
        if len(t.split()) > 8:
            ends = [h for h in page.search_for(words(t, 6, end=True)) if h.y0 >= r.y0 - 2]
            if ends: r |= ends[0]
        pos[n] = [pi + 1, [round(r.x0 / W, 4), round(r.y0 / H, 4), round(r.x1 / W, 4), round(r.y1 / H, 4)]]
        cur = pi
    note = f"{miss} units not found in the rendering (no highlight)" if miss else None
    return pos, note

def positions_job(args):
    st, n, src, kind, units = args
    cache = OUT / "cache" / st / f"{n}.bbox.json"
    h = sha(src)
    if cache.exists():
        c = json.loads(cache.read_text())
        if c.get("sha") == h: return st, n, c
    note, pos, pdf_rel = None, None, None
    try:
        if kind == "pdf":
            pos, note = pdf_positions(src, units)
        elif kind == "docx":
            pdf = OUT / "pdf" / st / f"{n}.pdf"; pdf.parent.mkdir(parents=True, exist_ok=True)
            if src.name in NO_RENDER:
                raise RuntimeError("LibreOffice cannot render this file (it hangs); showing its paragraphs and tables instead")
            with tempfile.TemporaryDirectory() as td:
                tmp = Path(td) / ("doc" + src.suffix)
                shutil.copy(src, tmp)
                subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{td}/profile", "--convert-to", "pdf", "--outdir", td, str(tmp)],
                               check=True, capture_output=True, timeout=600)
                shutil.copy(Path(td) / "doc.pdf", pdf)
            pos, note = docx_positions(pdf, units)
            pdf_rel = pdf.relative_to(ROOT).as_posix()
    except Exception as e:
        return st, n, {"sha": h, "pos": None, "note": f"positions not available: {e}", "pdf": None}   # not cached: retried next build
    c = {"sha": h, "pos": pos, "note": note, "pdf": pdf_rel}
    dump(cache, c)
    return st, n, c

def docx_blocks(src, units):
    """Fallback view of a .docx that LibreOffice cannot render: its paragraphs and tables in order, each mapped to
    its unit (same walk as extract_units.docx_units; used only if it reproduces the stored unit texts)."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from extract_units import norm
    d = docx.Document(str(src))
    blocks, n, texts = [], 0, []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d); t = norm(p.text)
            if t:
                n += 1; texts.append(t)
                st = (p.style.name if p.style is not None else "") or ""
                blocks.append({"n": n, "s": t, "h": st.lower().startswith(("heading", "title"))})
        elif tag == "tbl":
            rows = []
            for row in Table(el, d).rows:
                seen, cells = set(), []
                for cell in row.cells:
                    if id(cell._tc) in seen: continue
                    seen.add(id(cell._tc))
                    t = norm("\n".join(p.text for p in cell.paragraphs))
                    if t: n += 1; texts.append(t); cells.append([n, t])
                    else: cells.append([None, ""])
                rows.append(cells)
            blocks.append({"table": rows})
    stored = [u["text"] for u in units if u.get("style") != "textbox"]
    return blocks if texts == stored else None

# ---------- spreadsheet grid (the "original" view of .xlsx / .csv) ----------
def grid(src):
    if src.suffix.lower() == ".csv":
        with open(src, encoding="utf-8-sig", errors="ignore", newline="") as f:
            return [{"name": src.name, "rows": [[i, r] for i, r in enumerate(csv.reader(f), start=1) if any(c.strip() for c in r)]}]
    import openpyxl
    wb = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = ["" if c is None else str(c) for c in r]
            while vals and not vals[-1].strip(): vals.pop()
            if any(v.strip() for v in vals): rows.append([i, vals])
        out.append({"name": ws.title, "rows": rows})
    return out

# ---------- immigration-topic tags (from scripts/topic_tags.py) ----------
ANALYSIS = ROOT / "data" / "analysis"
METHODS = {"meaning": 0, "name": 1, "both": 2}

def load_topic_tags(index):
    """Write site/data/topics/ (the analysis: standards text only) and return {(state, file): [tag rows for the
    viewer, all categories]}."""
    tf = ANALYSIS / "topic_tags.csv"
    if not tf.exists():
        print("no data/analysis/topic_tags.csv: run scripts/topic_tags.py on main first; Topics and Charts tabs will be empty"); return {}
    meta = json.loads((ANALYSIS / "topic_tags_meta.json").read_text())
    topics = []
    for r in csv.DictReader(open(ANALYSIS / "topics.csv", encoding="utf-8")):
        topics.append({"no": int(r["topic_no"]), "era": r["era"], "topic": r["topic"], "years": r["years"], "description": r["description"],
                       "terms": [x.strip() for x in r["search_terms"].split(";") if x.strip()]})
    eras = list(dict.fromkeys(t["era"] for t in topics))
    states = sorted(index["states"])
    sidx = {s: i for i, s in enumerate(states)}
    fileno = {(s, f["name"]): f["n"] for s, v in index["states"].items() for f in v["files"]}
    acats = set(meta.get("analysis_categories", ["standards"]))
    words = [r for r in csv.DictReader(open(ANALYSIS / "text_words.csv", encoding="utf-8")) if r["category"] in acats]
    all_tags = list(csv.DictReader(open(tf, encoding="utf-8")))
    tags = [r for r in all_tags if r["category"] in acats]
    kf = ANALYSIS / "keyword_tags.csv"   # scripts/keyword_tags.py: the separate keyword counts
    kw = [r for r in csv.DictReader(open(kf, encoding="utf-8")) if r["category"] in acats] if kf.exists() else []
    grades = sorted({r["grades"] for r in words} | {r["grades"] for r in tags} | {r["grades"] for r in kw})
    gidx = {g: i for i, g in enumerate(grades)}
    courses = sorted({r["course"] for r in tags + kw if r["course"]})
    cidx = {c: i for i, c in enumerate(courses)}
    cat = {"standards": 0, "examples": 1, "backup": 2}
    # tag columns: x = index of the sentence in text.json, ms/ml = where the matched words sit in it, wp = word position in the document
    cols = {k: [] for k in ("st", "d", "u", "s", "c", "g", "t", "sc", "ts", "m", "co", "x", "ms", "ml", "wp")}
    text, tidx, per_doc = [], {}, {}
    for r in tags:
        d = fileno.get((r["state"], r["file"]), -1)
        sent = r["sentence"]
        if sent not in tidx: tidx[sent] = len(text); text.append(sent)
        ms = sent.find(r["matched"]) if r.get("matched") else -1
        vals = dict(st=sidx[r["state"]], d=d, u=int(r["unit"]), s=int(r["sentence_no"]), c=cat[r["category"]], g=gidx[r["grades"]], t=int(r["topic_no"]),
                    sc=round(float(r["score"]) * 1000), ts=round(float(r["top_score"]) * 1000), m=METHODS[r["method"]], co=cidx.get(r["course"], -1),
                    x=tidx[sent], ms=ms, ml=len(r["matched"]) if ms >= 0 else 0, wp=int(r["word_pos"]))
        for k, v in vals.items(): cols[k].append(v)
    # keyword matches: per keyword, the same columns, as one pseudo-topic (t=0, -1, ...) found by name with a full
    # score, and their own sentences
    kws = list(dict.fromkeys(r["keyword"] for r in kw))
    kdata = {}
    for ki, name in enumerate(kws):
        kcols, ktext, ktidx = {k: [] for k in cols}, [], {}
        for r in kw:
            if r["keyword"] != name: continue
            sent = r["sentence"]
            if sent not in ktidx: ktidx[sent] = len(ktext); ktext.append(sent)
            vals = dict(st=sidx[r["state"]], d=fileno.get((r["state"], r["file"]), -1), u=int(r["unit"]), s=int(r["sentence_no"]), c=cat[r["category"]],
                        g=gidx[r["grades"]], t=-ki, sc=1000, ts=1000, m=METHODS["name"], co=cidx.get(r["course"], -1), x=ktidx[sent], ms=int(r["sentence_char"]),
                        ml=len(r["matched"]), wp=int(r["word_pos"]))
            for k, v in vals.items(): kcols[k].append(v)
        kdata[name] = {"tags": kcols, "text": ktext}
    for r in all_tags:
        cs = int(r["char_start"]) if r.get("char_start") not in (None, "") else -1
        ce = int(r["char_end"]) if r.get("char_end") not in (None, "") else -1
        per_doc.setdefault((r["state"], r["file"]), []).append([int(r["unit"]), int(r["sentence_no"]), int(r["topic_no"]), round(float(r["score"]) * 1000),
                                                               round(float(r["top_score"]) * 1000), METHODS[r["method"]], cat[r["category"]], cs, ce, r["matched"]])
    wcols = {"st": [sidx[r["state"]] for r in words], "c": [cat[r["category"]] for r in words], "g": [gidx[r["grades"]] for r in words],
             "w": [int(r["words"]) for r in words], "ns": [int(r["sentences"]) for r in words]}
    T = OUT / "topics"
    dump(T / "list.json", {"meta": meta, "topics": topics, "eras": eras, "states": states, "grades": grades, "courses": courses})
    dump(T / "tags.json", cols)
    dump(T / "text.json", text)
    dump(T / "words.json", wcols)
    if kf.exists(): dump(T / "keyword.json", {"keywords": kdata})
    # database rows -> the tags in their standards text (indices into tags.json / keyword.json), for the Database tab
    def by_row(c):
        at = {}
        for i, (st, d, u) in enumerate(zip(c["st"], c["d"], c["u"])): at.setdefault((states[st], d, u), []).append(i)
        out = {}
        for st in states:
            rf = OUT / "rows" / f"{st}.json"
            if not rf.exists(): continue
            for r in json.loads(rf.read_text()):
                idx = [i for fi, u, _ in r["seg"]["body_standards"] for i in at.get((st, fi, u), [])]
                if idx: out[r["id"]] = sorted(set(idx))
        return out
    dump(T / "rowtags.json", {"tags": by_row(cols), "kw": {name: by_row(d["tags"]) for name, d in kdata.items()}})
    for name, d in kdata.items(): print(f"keyword {name}: {len(d['tags']['t']):,} matches in the analysis")
    print(f"topics: {len(tags):,} tags in the analysis ({'/'.join(sorted(acats))}), {len(all_tags):,} in the viewer, {len(grades)} grade groups")
    return per_doc

def set_parts(index):
    """Mark main documents and their subfiles (scripts/doc_parts.py) on the files of each state in index.json."""
    for st, s in index["states"].items():
        n = {f["name"]: f["n"] for f in s["files"]}
        for f in s["files"]: f.pop("subfiles", None); f.pop("part_of", None)
        for main, subs in doc_parts.parts(st).items():
            if main not in n: continue
            s["files"][n[main]]["subfiles"] = [n[x] for x in subs if x in n]
            for x in subs:
                if x in n: s["files"][n[x]]["part_of"] = n[main]

def topics_only():
    idx_path = OUT / "index.json"
    index = json.loads(idx_path.read_text())
    set_parts(index)
    per_doc = load_topic_tags(index)
    for st, s in index["states"].items():
        for f in s["files"]:
            p = OUT / "docs" / st / f"{f['n']}.json"
            if f["status"] != "use" or not p.exists(): continue
            doc = json.loads(p.read_text())
            doc["tags"] = per_doc.get((st, doc["file"]), [])
            dump(p, doc)
    dump(idx_path, index)
    print("wrote", OUT / "topics", "and the tags of every document")

def main():
    if "--topics" in sys.argv: return topics_only()
    cfg, names = load_cfg()
    states = sys.argv[1:] or sorted(cfg)
    man = {(r["abbr"], Path(r["local_path"]).name): r for r in csv.DictReader(open(ROOT / "data" / "manifest.csv"))}
    idx_path = OUT / "index.json"
    index = json.loads(idx_path.read_text()) if idx_path.exists() and sys.argv[1:] else {"states": {}}
    jobs, docs = [], {}
    for st in states:
        c = cfg[st]
        files = []
        for i, (f, v) in enumerate(c.get("files", {}).items()):
            m = man.get((st, f), {})
            files.append({"n": i, "name": f, "status": v, "path": m.get("local_path", ""), "url": m.get("final_url") or m.get("source_url") or "",
                          "type": Path(f).suffix.lower().lstrip(".")})
        fileno = {f["name"]: f["n"] for f in files}
        if not files:
            rows = [{"id": f"{st}_NONE", "state": st, "grade": "[]", "course": "", "course_category": "", "course_title": "", "source_files": "",
                     "issues": "; ".join(c.get("issues", ["no source file"])), **{k: "" for k in COLS}}]
        else:
            rows = assemble_state(st, cfg, names)
        # unit -> [(row id, column, kept)]
        member = {}
        for r in rows:
            seg = {}
            for col in COLS:
                items = r.get("_prov", {}).get(col, [])
                seg[col] = [[fileno.get(f, -1), u, t] for f, u, t, kept in items if kept]
                assert "\n".join(s[2] for s in seg[col]) == r[col], (r["id"], col)
                for f, u, t, kept in items:
                    member.setdefault((f, u), []).append([r["id"], col[5:] if col.startswith("body_") else col, int(kept)])
            r.pop("_prov", None); r["seg"] = seg
            r["words"] = [len(r.pop(k).split()) for k in COLS]   # the text itself is in seg (checked equal above)
        dump(OUT / "rows" / f"{st}.json", rows)
        index["states"][st] = {"name": STATE_NAMES.get(st, st), "files": files, "issues": c.get("issues", []),
            "rows": [{"id": r["id"], "grade": r["grade"], "course": r["course"], "course_category": r["course_category"], "course_title": r["course_title"],
                      "words": r["words"], "files": [fileno[f] for f in r["source_files"].split("; ") if f in fileno]} for r in rows]}
        for fi in files:
            if fi["status"] != "use": continue
            uf = CLEAN / "units" / st / f"{fi['name']}.units.jsonl"
            lf = CLEAN / "labels" / st / f"{fi['name']}.labels.json"
            if not uf.exists() or not lf.exists(): continue
            units = [json.loads(l) for l in open(uf)]
            lab = json.loads(lf.read_text())
            labels = lab.get("labels", [])
            which = {}
            for li, l in enumerate(labels):
                for x in parse_units(l["units"]): which.setdefault(x, li)
            allrow = f"{st}_ALL"
            doc = {"state": st, "n": fi["n"], "file": fi["name"], "type": fi["type"], "path": fi["path"], "url": fi["url"],
                   "structure": lab.get("structure", ""), "issues": lab.get("issues", []),
                   "labels": [{"units": l["units"], "grades": l["grades"], "course": l.get("course"), "cat": l["cat"], "section": l.get("section", ""),
                               "keep": bool(l.get("keep_in_rows"))} for l in labels],
                   "units": []}
            for u in units:
                n = int(u["id"][1:])
                rws = member.get((fi["name"], n), [])
                if u.get("furniture"): rws = [[allrow, "backup", 1]]
                doc["units"].append({"n": n, "p": u["page"], "k": u["kind"], "t": unit_tag(u), "x": u.get("x"), "s": u["text"],
                                     "f": int(bool(u.get("furniture"))), "l": which.get(n, -1), "r": rws})
            src = ROOT / fi["path"]
            if fi["type"] in ("xlsx", "csv") and src.exists():
                doc["grid"] = grid(src)
            docs[(st, fi["n"])] = doc
            if fi["type"] in ("pdf", "docx") and src.exists():
                jobs.append((st, fi["n"], src, fi["type"], units))
        print(st, len(rows), "rows", sum(f["status"] == "use" for f in files), "docs", flush=True)
    with ProcessPoolExecutor() as ex:
        for st, n, c in ex.map(positions_job, jobs):
            doc = docs[(st, n)]
            pos = {int(k): v for k, v in (c["pos"] or {}).items()}
            for u in doc["units"]:
                if u["n"] in pos:
                    u["p"], u["b"] = pos[u["n"]]
            doc["view"] = c["pdf"] or doc["path"]
            if doc["type"] == "docx" and not c["pdf"]:
                blocks = docx_blocks(ROOT / doc["path"], [json.loads(l) for l in open(CLEAN / "units" / st / f"{doc['file']}.units.jsonl")])
                if blocks:
                    doc["blocks"] = blocks
                    doc["view_note"] = "LibreOffice cannot render this Word file, so it is shown as its own paragraphs and tables, read directly from the .docx."
            if c["note"]: doc["view_note"] = c["note"]
            hit = sum("b" in u for u in doc["units"])
            print(f"  {st} {doc['file'][:60]}: {hit}/{len(doc['units'])} units placed" + (f" ({c['note']})" if c["note"] else ""), flush=True)
    per_doc = load_topic_tags(index)
    for (st, n), doc in docs.items():
        doc.setdefault("view", doc["path"])
        doc["tags"] = per_doc.get((st, doc["file"]), [])
        dump(OUT / "docs" / st / f"{n}.json", doc)
    index["states"] = dict(sorted(index["states"].items()))
    set_parts(index)
    index["totals"] = {"states": len(index["states"]), "rows": sum(len(s["rows"]) for s in index["states"].values()),
                       "files": sum(len(s["files"]) for s in index["states"].values()),
                       "used": sum(f["status"] == "use" for s in index["states"].values() for f in s["files"])}
    dump(idx_path, index)
    print("wrote", OUT)

if __name__ == "__main__":
    main()
