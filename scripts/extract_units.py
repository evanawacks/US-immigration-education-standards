#!/usr/bin/env python3
"""Clean extraction layer: split each source document into ordered 'units'
(paragraphs, table cells, spreadsheet rows) with stable IDs.

Output per document in data/clean/units/<ABBR>/:
  <stem>.units.jsonl  one unit per line: {id, page, kind, table, row, col, x, text, furniture}
  <stem>.view.txt     compact view for reading/labeling:  'u12 text' / 'u13 [T1 R2 C3] text'
Repeated page furniture (running headers/footers/page numbers) is detected and marked
furniture=true; it is omitted from the view and goes to the backup column automatically.

Usage: python scripts/extract_units.py [ABBR ...]
"""
import csv, json, re, sys, collections
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "clean" / "units"

def norm(s):
    s = s.replace("­", "").replace("ﬁ", "fi").replace("ﬂ", "fl")
    return re.sub(r"[ \t\r\f\v]+", " ", re.sub(r"\s*\n\s*", " ", s)).strip()

# ---------------- PDF ----------------
def pdf_units(path):
    import pymupdf
    doc = pymupdf.open(str(path))
    units = []
    for pno, page in enumerate(doc, start=1):
        try:
            tabs = page.find_tables().tables
        except Exception:
            tabs = []
        tboxes = [pymupdf.Rect(t.bbox) for t in tabs]
        items = []  # (y0, order, unit)
        # tables -> one unit per cell
        for ti, t in enumerate(tabs, start=1):
            try:
                rows = t.extract()
            except Exception:
                continue
            k = 0
            for ri, row in enumerate(rows, start=1):
                for ci, cell in enumerate(row, start=1):
                    if cell and norm(cell):
                        items.append((t.bbox[1], k, {"page": pno, "kind": "cell", "table": ti, "row": ri, "col": ci, "x": None, "text": norm(cell)}))
                        k += 1
        # text outside tables -> lines grouped into paragraph units
        d = page.get_text("dict")
        order = 0
        for b in d["blocks"]:
            if b.get("type") != 0: continue
            cur, cur_x, cur_y1 = [], None, None
            def flush():
                nonlocal cur
                if cur:
                    txt = norm(" ".join(c[2] for c in cur))
                    if txt:
                        items.append((cur[0][1], 10000 + len(items), {"page": pno, "kind": "text", "table": None, "row": None, "col": None, "x": round(cur[0][0]), "y1": cur[-1][3], "y0": cur[0][1], "text": txt}))
                cur = []
            for ln in b["lines"]:
                txt = "".join(sp["text"] for sp in ln["spans"])
                if not txt.strip(): continue
                x0, y0, x1, y1 = ln["bbox"]
                c = pymupdf.Rect(ln["bbox"])
                mid = pymupdf.Point((x0 + x1) / 2, (y0 + y1) / 2)
                if any(r.contains(mid) for r in tboxes): continue
                h = max(1.0, y1 - y0)
                if cur and (abs(x0 - cur_x) > 60 or y0 - cur_y1 > 1.2 * h or y0 < cur[-1][1] - 2):
                    flush()
                if not cur: cur_x = x0
                cur.append((x0, y0, txt, y1)); cur_y1 = y1
            flush()
        # stable order: by vertical position of block start, keeping stream order inside
        items.sort(key=lambda it: (round(it[0] / 4), it[1]))
        # mark side-by-side text units (need x to tell columns apart)
        txt_units = [u for _, _, u in items if u["kind"] == "text"]
        side = [any(v is not u and abs(v["x"] - u["x"]) > 60 and v["y0"] < u["y1"] and u["y0"] < v["y1"] for v in txt_units) for u in txt_units]
        for u, sd in zip(txt_units, side):
            if not sd: u["x"] = None
        for _, _, u in items:
            u.pop("y0", None); u.pop("y1", None)
            units.append(u)
    return units, len(doc)

# ---------------- DOCX ----------------
def docx_units(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    units, tn = [], 0
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d)
            t = norm(p.text)
            if t:
                st = (p.style.name if p.style is not None else "") or ""
                units.append({"page": None, "kind": "text", "table": None, "row": None, "col": None, "x": None, "style": st if st.lower().startswith(("heading", "title")) else None, "text": t})
        elif tag == "tbl":
            tn += 1
            tb = Table(el, d)
            for ri, row in enumerate(tb.rows, start=1):
                seen = set()
                for ci, cell in enumerate(row.cells, start=1):
                    if id(cell._tc) in seen: continue
                    seen.add(id(cell._tc))
                    t = norm("\n".join(p.text for p in cell.paragraphs))
                    if t:
                        units.append({"page": None, "kind": "cell", "table": tn, "row": ri, "col": ci, "x": None, "text": t})
    # text boxes are skipped by python-docx; append their paragraphs at the end (keeps earlier unit IDs stable)
    have = {u["text"] for u in units}
    for tb in d.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent"):
        for p in tb.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
            t = norm("".join(x.text or "" for x in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")))
            if t and t not in have:
                have.add(t)
                units.append({"page": None, "kind": "text", "table": None, "row": None, "col": None, "x": None, "style": "textbox", "text": t})
    return units, None

# ---------------- XLSX / CSV ----------------
def xlsx_units(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    units = []
    for si, ws in enumerate(wb.worksheets, start=1):
        units.append({"page": si, "kind": "sheet", "table": si, "row": 0, "col": None, "x": None, "text": f"SHEET: {ws.title}"})
        for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [norm(str(c)) for c in row if c is not None and norm(str(c))]
            if vals:
                units.append({"page": si, "kind": "row", "table": si, "row": ri, "col": None, "x": None, "text": " | ".join(vals)})
    return units, len(wb.worksheets)

def csv_units(path):
    units = []
    with open(path, encoding="utf-8-sig", errors="ignore", newline="") as f:
        for ri, row in enumerate(csv.reader(f), start=1):
            vals = [norm(c) for c in row if norm(c)]
            if vals:
                units.append({"page": None, "kind": "row", "table": 1, "row": ri, "col": None, "x": None, "text": " | ".join(vals)})
    return units, None

def mark_furniture(units, npages):
    if not npages or npages < 4: return
    pages_by = collections.defaultdict(set)
    key = lambda t: re.sub(r"\d+", "#", t.lower())[:120]
    for u in units:
        if u["kind"] == "text" and len(u["text"]) < 160: pages_by[key(u["text"])].add(u["page"])
    rep = {k for k, v in pages_by.items() if len(v) >= max(3, 0.3 * npages)}
    for u in units:
        u["furniture"] = bool(u["kind"] == "text" and len(u["text"]) < 160 and (key(u["text"]) in rep or re.fullmatch(r"(page )?\d{1,3}( of \d{1,3})?|[ivx]{1,5}", u["text"].lower())))

def render(units):
    hdr = collections.defaultdict(list)
    for u in units:
        if u.get("furniture") and u["page"] is not None:
            t = u["text"][:70]
            if t not in hdr[u["page"]] and not re.fullmatch(r"(page )?\d{1,3}( of \d{1,3})?|[ivx]{1,5}", t.lower()):
                hdr[u["page"]].append(t)
    out, page = [], object()
    for u in units:
        if u["page"] != page and u["kind"] in ("text", "cell"):
            page = u["page"]
            if page is not None:
                h = " | ".join(hdr.get(page, [])[:5])
                out.append(f"=== page {page} ===" + (f" [running header: {h}]" if h else ""))
        if u.get("furniture"): continue
        tag = ""
        if u["kind"] == "cell": tag = f"[T{u['table']} R{u['row']} C{u['col']}] "
        elif u["kind"] == "row": tag = f"[R{u['row']}] "
        elif u.get("style"): tag = f"[{u['style']}] "
        if u.get("x") is not None: tag += f"(x={u['x']}) "
        out.append(f"{u['id']} {tag}{u['text']}")
    return "\n".join(out) + "\n"

def run(abbrs):
    rows = list(csv.DictReader(open(ROOT / "data" / "manifest.csv")))
    stats = []
    for r in rows:
        if abbrs and r["abbr"] not in abbrs: continue
        p = ROOT / r["local_path"]; ext = p.suffix.lower()
        fn = {".pdf": pdf_units, ".docx": docx_units, ".xlsx": xlsx_units, ".csv": csv_units}.get(ext)
        if not fn: continue
        try:
            units, npages = fn(p)
        except Exception as e:
            print("ERROR", p, e); continue
        for i, u in enumerate(units, start=1): u["id"] = f"u{i}"
        mark_furniture(units, npages)
        d = OUT / r["abbr"]; d.mkdir(parents=True, exist_ok=True)
        stem = p.name
        with open(d / f"{stem}.units.jsonl", "w") as f:
            for u in units: f.write(json.dumps(u, ensure_ascii=False) + "\n")
        view = render(units)
        (d / f"{stem}.view.txt").write_text(view)
        stats.append((r["abbr"], stem, len(units), sum(u.get("furniture", False) for u in units), len(view.split())))
        print(r["abbr"], stem[:60], "units", len(units), "furniture", stats[-1][3], "view words", stats[-1][4], flush=True)

if __name__ == "__main__":
    run(set(sys.argv[1:]))
